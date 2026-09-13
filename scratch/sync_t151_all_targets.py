import sqlite3
import json
import os
import openpyxl

with open('scratch/t151_50_built_questions.json', 'r', encoding='utf-8') as f:
    built_qs = json.load(f)

print(f"Loaded {len(built_qs)} built questions for Topic 151.")

# Get existing IDs for Topic 151 from server/etuition.db
conn_main = sqlite3.connect('server/etuition.db')
cursor_main = conn_main.cursor()
cursor_main.execute("SELECT id FROM questions WHERE topic_id = 151 ORDER BY id")
existing_ids = [r[0] for r in cursor_main.fetchall()]
conn_main.close()

print(f"Existing IDs for Topic 151 in server/etuition.db: {len(existing_ids)} (e.g., {existing_ids[:5]})")

# Map 50 questions to existing IDs
full_records = []
for idx, bq in enumerate(built_qs):
    q_id = existing_ids[idx] if idx < len(existing_ids) else (134433 + idx)
    img_url = f"/images/g9_t151_q{bq['num']}.svg"
    opts_json = json.dumps(bq['options'], ensure_ascii=False)
    steps_json = json.dumps(bq['steps'], ensure_ascii=False)
    
    record = {
        'id': q_id,
        'topic_id': 151,
        'question_title': f"Q{bq['num']}: {bq['title']}",
        'question_text': bq['text'],
        'math_formula': bq['formula'],
        'options_json': opts_json,
        'correct_answer': bq['answer'],
        'hint': bq['hint'],
        'working_steps_json': steps_json,
        'created_by': 1,
        'question_type': 'multiple_choice',
        'image_url': img_url,
        'image_alt': bq['img_title'],
        'difficulty': 2,
        'show_image': 1,
        'show_formula': 1
    }
    full_records.append(record)

# 1. Update all 3 SQLite databases
dbs = ['server/etuition.db', 'QBank/server/qbank.db', 'QBank/server/etuition.db']
for db_path in dbs:
    if os.path.exists(db_path):
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        cursor.execute("DELETE FROM questions WHERE topic_id = 151")
        
        for r in full_records:
            cursor.execute("""
                INSERT INTO questions (
                    id, topic_id, question_title, question_text, math_formula,
                    options_json, correct_answer, hint, working_steps_json,
                    image_url, image_alt, difficulty, show_image, show_formula
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                r['id'], r['topic_id'], r['question_title'], r['question_text'], r['math_formula'],
                r['options_json'], r['correct_answer'], r['hint'], r['working_steps_json'],
                r['image_url'], r['image_alt'], r['difficulty'], r['show_image'], r['show_formula']
            ))
        conn.commit()
        conn.close()
        print(f"Updated SQLite DB: {db_path}")

# 2. Update all 4 JSON files
json_paths = [
    'questions.json',
    'QBank/questions.json',
    'QBank/public/questions.json',
    'QBank/dist/questions.json'
]

for jp in json_paths:
    if os.path.exists(jp):
        with open(jp, 'r', encoding='utf-8') as f:
            data = json.load(f)
            
        data = [q for q in data if str(q.get('topic_id')) != '151']
        for r in full_records:
            data.append(r)
            
        with open(jp, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f"Updated JSON: {jp}")

# 3. Update Excel bank
excel_path = 'QBank/questions_bank.xlsx'
if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    
    col_map = {h: idx+1 for idx, h in enumerate(headers)}
    rec_map = {r['id']: r for r in full_records}
    
    for row_idx in range(2, sheet.max_row + 1):
        qid_val = sheet.cell(row=row_idx, column=1).value
        if qid_val in rec_map:
            r = rec_map[qid_val]
            opts = json.loads(r['options_json'])
            
            if 'Question Title' in col_map: sheet.cell(row=row_idx, column=col_map['Question Title'], value=r['question_title'])
            if 'Question Text' in col_map: sheet.cell(row=row_idx, column=col_map['Question Text'], value=r['question_text'])
            if 'LaTeX Formula Expression' in col_map: sheet.cell(row=row_idx, column=col_map['LaTeX Formula Expression'], value=r['math_formula'])
            if 'Option A' in col_map and len(opts)>0: sheet.cell(row=row_idx, column=col_map['Option A'], value=opts[0])
            if 'Option B' in col_map and len(opts)>1: sheet.cell(row=row_idx, column=col_map['Option B'], value=opts[1])
            if 'Option C' in col_map and len(opts)>2: sheet.cell(row=row_idx, column=col_map['Option C'], value=opts[2])
            if 'Option D' in col_map and len(opts)>3: sheet.cell(row=row_idx, column=col_map['Option D'], value=opts[3])
            if 'Correct Answer' in col_map: sheet.cell(row=row_idx, column=col_map['Correct Answer'], value=r['correct_answer'])
            if 'Hint' in col_map: sheet.cell(row=row_idx, column=col_map['Hint'], value=r['hint'])
            if 'Working Steps' in col_map: sheet.cell(row=row_idx, column=col_map['Working Steps'], value=r['working_steps_json'])
            if 'Image URL' in col_map: sheet.cell(row=row_idx, column=col_map['Image URL'], value=r['image_url'])
            if 'Image ALT' in col_map: sheet.cell(row=row_idx, column=col_map['Image ALT'], value=r['image_alt'])

    wb.save(excel_path)
    print(f"Updated Excel Bank: {excel_path}")

print("Synchronization for Topic 151 complete across all DBs, JSONs, and Excel!")
