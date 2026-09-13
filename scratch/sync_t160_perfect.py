import sqlite3
import json
import os
import openpyxl

with open("scratch/t160_50_perfect_built.json", "r", encoding="utf-8") as f:
    t160_data = json.load(f)

print(f"Loaded {len(t160_data)} perfect questions for T160 sync.")

# 1. Update SQLite Databases
dbs = [
    "server/etuition.db",
    "QBank/server/qbank.db",
    "QBank/server/etuition.db"
]

for db_path in dbs:
    if not os.path.exists(db_path):
        continue
    
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    cursor.execute("SELECT id FROM questions WHERE topic_id = 160 ORDER BY id ASC")
    existing_rows = cursor.fetchall()
    
    for idx, q_info in enumerate(t160_data):
        if idx < len(existing_rows):
            q_id = existing_rows[idx][0]
            opts_json = json.dumps(q_info["options"], ensure_ascii=False)
            cursor.execute("""
                UPDATE questions 
                SET question_title = ?,
                    question_text = ?,
                    math_formula = ?,
                    options_json = ?,
                    correct_answer = ?,
                    show_image = ?,
                    image_url = ?
                WHERE id = ?
            """, (
                q_info["title"],
                q_info["text"],
                q_info["formula"],
                opts_json,
                q_info["correct"],
                q_info["show_image"],
                q_info["image_url"],
                q_id
            ))
    
    conn.commit()
    conn.close()
    print(f"Successfully updated SQLite DB: {db_path}")

# 2. Update JSON Datastores
json_paths = [
    "questions.json",
    "QBank/questions.json",
    "QBank/public/questions.json",
    "QBank/dist/questions.json"
]

for jpath in json_paths:
    if not os.path.exists(jpath):
        continue
    
    with open(jpath, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    t160_items = []
    if isinstance(data, list):
        t160_items = [q for q in data if q.get("topic_id") == 160]
    elif isinstance(data, dict) and "questions" in data:
        t160_items = [q for q in data["questions"] if q.get("topic_id") == 160]
    
    t160_items.sort(key=lambda x: x.get("id", 0))
    for idx, q_info in enumerate(t160_data):
        if idx < len(t160_items):
            target = t160_items[idx]
            target["question_title"] = q_info["title"]
            target["question_text"] = q_info["text"]
            target["math_formula"] = q_info["formula"]
            target["options_json"] = json.dumps(q_info["options"], ensure_ascii=False)
            target["options"] = q_info["options"]
            target["correct_answer"] = q_info["correct"]
            target["correct_option"] = q_info["correct"]
            target["show_image"] = q_info["show_image"]
            target["image_url"] = q_info["image_url"]
    
    with open(jpath, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"Successfully updated JSON datastore: {jpath}")

# 3. Update Excel Question Bank
excel_path = "QBank/questions_bank.xlsx"
if os.path.exists(excel_path):
    try:
        wb = openpyxl.load_workbook(excel_path)
        sheet = wb.active
        
        t_col = 5
        title_col = 8
        text_col = 9
        formula_col = 10
        opta_c = 12
        optb_c = 13
        optc_c = 14
        optd_c = 15
        corr_c = 16
        img_c = 19
        
        matching_rows = []
        for r in range(2, sheet.max_row + 1):
            cval = sheet.cell(row=r, column=t_col).value
            if cval is not None:
                try:
                    if int(cval) == 160:
                        matching_rows.append(r)
                except ValueError:
                    pass
        matching_rows.sort()
        
        for idx, q_info in enumerate(t160_data):
            if idx < len(matching_rows):
                r = matching_rows[idx]
                sheet.cell(row=r, column=title_col, value=q_info["title"])
                sheet.cell(row=r, column=text_col, value=q_info["text"])
                sheet.cell(row=r, column=formula_col, value=q_info["formula"])
                
                opts = q_info["options"]
                sheet.cell(row=r, column=opta_c, value=opts[0] if len(opts)>0 else "")
                sheet.cell(row=r, column=optb_c, value=opts[1] if len(opts)>1 else "")
                sheet.cell(row=r, column=optc_c, value=opts[2] if len(opts)>2 else "")
                sheet.cell(row=r, column=optd_c, value=opts[3] if len(opts)>3 else "")
                
                sheet.cell(row=r, column=corr_c, value=q_info["correct"])
                sheet.cell(row=r, column=img_c, value=q_info["image_url"] if q_info["show_image"] else "")
        
        wb.save(excel_path)
        print(f"Successfully updated Excel bank: {excel_path}")
    except Exception as e:
        print(f"Error updating Excel bank: {e}")

print("Synchronization complete!")
