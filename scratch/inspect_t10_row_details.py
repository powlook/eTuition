import openpyxl

wb = openpyxl.load_workbook('QBank/resources/Grade_9_Math_50_Unique_Questions.xlsx')

for sheet_name in ['All 850 Questions', 'Measurement & Geometry (MG)']:
    ws = wb[sheet_name]
    print(f"\n=== SHEET: {sheet_name} ===")
    headers = [str(ws.cell(row=1, column=c).value).encode('ascii', 'ignore').decode('ascii') for c in range(1, ws.max_column+1)]
    print("Headers:", headers)
    for r in range(2, ws.max_row+1):
        v1 = str(ws.cell(row=r, column=1).value)
        if 'T10-Q001' in v1 or 'T10-Q003' in v1:
            print(f"\nRow {r} ({v1}):")
            for c in range(1, min(ws.max_column+1, 15)):
                val = ws.cell(row=r, column=c).value
                val_safe = str(val).encode('ascii', 'ignore').decode('ascii')
                print(f"  Col {c} ({headers[c-1]}): {val_safe[:100]}")
