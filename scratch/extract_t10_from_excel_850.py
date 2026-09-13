import openpyxl
import json

wb = openpyxl.load_workbook('QBank/resources/Grade_9_Math_50_Unique_Questions.xlsx')

print("Sheet names in workbook:", wb.sheetnames)

ws = wb['Measurement & Geometry (MG)']
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
print("MG Headers:", headers)

for r in range(1, ws.max_row+1):
    val = str(ws.cell(row=r, column=2).value)
    if 'T10' in val or 'Trigonometric' in val:
        print(f"Row {r}: Topic={val} | Col1={ws.cell(row=r, column=1).value} | Col3={ws.cell(row=r, column=3).value}")
