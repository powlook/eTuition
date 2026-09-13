import sqlite3
import openpyxl
import json

print("=== 1. SERVER/ETUITION.DB FOR TOPIC 152 ===")
conn = sqlite3.connect('server/etuition.db')
c = conn.cursor()
c.execute("SELECT id, question_title, question_text, math_formula, correct_answer, image_url, show_image FROM questions WHERE topic_id = 152 ORDER BY id")
db_qs = c.fetchall()
print(f"Total questions in server/etuition.db for Topic 152: {len(db_qs)}")
for idx, r in enumerate(db_qs):
    t_safe = r[1].encode('ascii', 'ignore').decode('ascii') if r[1] else ''
    txt_safe = r[2][:80].encode('ascii', 'ignore').decode('ascii') if r[2] else ''
    print(f"Q{idx+1} (ID {r[0]}): Title='{t_safe}' | Img={r[5]} | Text='{txt_safe}'")
conn.close()

print("\n=== 2. QBANK/QUESTIONS_BANK.XLSX FOR TOPIC 152 ===")
wb = openpyxl.load_workbook('QBank/questions_bank.xlsx')
ws = wb.active
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
col_map = {h: idx+1 for idx, h in enumerate(headers) if h}

excel_qs = []
for r in range(2, ws.max_row+1):
    topic_id_val = str(ws.cell(row=r, column=col_map.get('Topic ID', 2)).value)
    if topic_id_val == '152':
        qid = ws.cell(row=r, column=1).value
        q_title = ws.cell(row=r, column=col_map.get('Question Title', 3)).value
        q_text = ws.cell(row=r, column=col_map.get('Question Text', 4)).value
        img_url = ws.cell(row=r, column=col_map.get('Image URL', 13)).value
        excel_qs.append((qid, q_title, q_text, img_url))

print(f"Total questions in QBank/questions_bank.xlsx for Topic 152: {len(excel_qs)}")
for idx, eq in enumerate(excel_qs[:10]):
    t_safe = str(eq[1]).encode('ascii', 'ignore').decode('ascii')
    txt_safe = str(eq[2])[:80].encode('ascii', 'ignore').decode('ascii')
    print(f"Q{idx+1} (ID {eq[0]}): Title='{t_safe}' | Img={eq[3]} | Text='{txt_safe}'")
