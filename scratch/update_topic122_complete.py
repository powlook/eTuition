import sqlite3
import json
import os
import openpyxl
import random

db_paths = [
    'server/etuition.db',
    'QBank/server/qbank.db',
    'QBank/server/etuition.db'
]

json_paths = [
    'questions.json',
    'QBank/questions.json',
    'QBank/public/questions.json',
    'QBank/dist/questions.json'
]

excel_path = 'QBank/questions_bank.xlsx'

# 1. Fetch current topic 122 questions from server/etuition.db
conn = sqlite3.connect('server/etuition.db')
c = conn.cursor()
c.execute("""
    SELECT id, topic_id, question_title, question_text, math_formula, options_json, correct_answer, hint, working_steps_json, image_url, image_alt, difficulty, show_image, show_formula
    FROM questions
    WHERE topic_id = 122 OR question_title LIKE '%T122%'
    ORDER BY id ASC
""")
rows = c.fetchall()
conn.close()

print(f"Fetched {len(rows)} topic 122 questions from DB.")

updated_questions = []

for r in rows:
    qid, tid, title, text, formula, opts_json, ans, hint, steps_json, img, alt, diff, show_img, show_form = r
    opts = json.loads(opts_json) if opts_json else []
    steps = json.loads(steps_json) if steps_json else []
    
    # Specific fixes
    # Q128679 (Item 6): Stem 4 | 2, 5; Stem 5 | 1, 3, 8; Stem 6 | 0, 4, 7
    if qid == 128679 or '128679' in str(qid):
        ans = "Median score = 55.5"
        opts = [
            "Median score = 52.0",
            "Median score = 55.5",
            "Median score = 58.0",
            "Median score = 60.5"
        ]
        formula = r"\text{Median} = \text{Middle value when ordered } (N = 8 \implies \text{average of } 4^{\text{th}} \text{ and } 5^{\text{th}} \text{ values})"
        steps = [
            "1. List all 8 quiz scores in ascending order from the stem-and-leaf plot: 42, 45, 51, 53, 58, 60, 64, 67.",
            "2. Identify the 4th value (53) and 5th value (58).",
            "3. Calculate the median score: Median = (53 + 58) / 2 = 55.5."
        ]
        img = "/images/dp_t122_q6.svg"
        hint = "Order the values from the stem-and-leaf plot: 42, 45, 51, 53, 58, 60, 64, 67. The median is the average of 53 and 58."

    # Remove image URLs for purely conceptual text questions so they don't show generic/misleading diagrams
    conceptual_ids = [128675, 128677, 128685, 128687, 128695, 128697, 128705, 128707, 128715, 128717]
    if qid in conceptual_ids:
        img = ""

    # Ensure correct answer is present in options
    if ans not in opts:
        opts[0] = ans
        
    updated_questions.append({
        'id': qid,
        'topic_id': tid,
        'question_title': title,
        'question_text': text,
        'math_formula': formula,
        'options': opts,
        'correct_answer': ans,
        'hint': hint,
        'working_steps': steps,
        'image_url': img,
        'image_alt': alt,
        'difficulty': diff,
        'show_image': show_img,
        'show_formula': show_form
    })



# Explicitly balance answer positions across A, B, C, D
target_letters = ['A']*13 + ['B']*13 + ['C']*12 + ['D']*12
random.seed(42)
random.shuffle(target_letters)
pos_counts = {'A': 0, 'B': 0, 'C': 0, 'D': 0}

for idx_q, q in enumerate(updated_questions):
    ans = q['correct_answer']
    opts = list(q['options'])
    if ans not in opts:
        opts.append(ans)
    # Remove ans from distractor list
    distractors = [o for o in opts if o != ans]
    target_letter = target_letters[idx_q]
    target_idx = ord(target_letter) - 65
    
    # Construct 4 options with ans at target_idx
    new_opts = [None] * 4
    new_opts[target_idx] = ans
    
    d_idx = 0
    for i in range(4):
        if i != target_idx:
            if d_idx < len(distractors):
                new_opts[i] = distractors[d_idx]
                d_idx += 1
            else:
                new_opts[i] = f"Option {i+1}"
    q['options'] = new_opts
    pos_counts[target_letter] += 1


print(f"Option position distribution for Topic 122: {pos_counts}")

# 2. Update SQLite DBs
for db_path in db_paths:
    if not os.path.exists(db_path):
        continue
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    for q in updated_questions:
        c.execute("""
            UPDATE questions
            SET question_title = ?,
                question_text = ?,
                math_formula = ?,
                options_json = ?,
                correct_answer = ?,
                hint = ?,
                working_steps_json = ?,
                image_url = ?
            WHERE id = ?
        """, (
            q['question_title'],
            q['question_text'],
            q['math_formula'],
            json.dumps(q['options'], ensure_ascii=False),
            q['correct_answer'],
            q['hint'],
            json.dumps(q['working_steps'], ensure_ascii=False),
            q['image_url'],
            q['id']
        ))
    conn.commit()
    conn.close()
    print(f"Updated SQLite DB: {db_path}")

# 3. Update JSON files
for json_path in json_paths:
    if not os.path.exists(json_path):
        continue
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    # Check structure
    q_map = {q['id']: q for q in updated_questions}
    
    if isinstance(data, list):
        for item in data:
            if item.get('id') in q_map:
                u = q_map[item['id']]
                item['question_title'] = u['question_title']
                item['question_text'] = u['question_text']
                item['math_formula'] = u['math_formula']
                item['options'] = u['options']
                item['options_json'] = json.dumps(u['options'], ensure_ascii=False)
                item['correct_answer'] = u['correct_answer']
                item['hint'] = u['hint']
                item['working_steps'] = u['working_steps']
                item['working_steps_json'] = json.dumps(u['working_steps'], ensure_ascii=False)
                item['image_url'] = u['image_url']
    elif isinstance(data, dict) and 'questions' in data:
        for item in data['questions']:
            if item.get('id') in q_map:
                u = q_map[item['id']]
                item['question_title'] = u['question_title']
                item['question_text'] = u['question_text']
                item['math_formula'] = u['math_formula']
                item['options'] = u['options']
                item['options_json'] = json.dumps(u['options'], ensure_ascii=False)
                item['correct_answer'] = u['correct_answer']
                item['hint'] = u['hint']
                item['working_steps'] = u['working_steps']
                item['working_steps_json'] = json.dumps(u['working_steps'], ensure_ascii=False)
                item['image_url'] = u['image_url']

    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"Updated JSON file: {json_path}")

# 4. Update Excel file if exists
if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    # Find headers
    headers = [cell.value for cell in ws[1]]
    id_col = headers.index('id') if 'id' in headers else -1
    q_map = {q['id']: q for q in updated_questions}
    
    if id_col != -1:
        for row in ws.iter_rows(min_row=2):
            val = row[id_col].value
            if val in q_map:
                u = q_map[val]
                for col_idx, col_name in enumerate(headers):
                    if col_name == 'question_title':
                        row[col_idx].value = u['question_title']
                    elif col_name == 'question_text':
                        row[col_idx].value = u['question_text']
                    elif col_name == 'math_formula':
                        row[col_idx].value = u['math_formula']
                    elif col_name in ('options', 'options_json'):
                        row[col_idx].value = json.dumps(u['options'], ensure_ascii=False)
                    elif col_name == 'correct_answer':
                        row[col_idx].value = u['correct_answer']
                    elif col_name == 'hint':
                        row[col_idx].value = u['hint']
                    elif col_name in ('working_steps', 'working_steps_json'):
                        row[col_idx].value = json.dumps(u['working_steps'], ensure_ascii=False)
                    elif col_name == 'image_url':
                        row[col_idx].value = u['image_url']
        wb.save(excel_path)
        print(f"Updated Excel file: {excel_path}")

print("Topic 122 update complete!")
