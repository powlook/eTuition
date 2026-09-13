import openpyxl

wb = openpyxl.load_workbook('QBank/resources/Grade_8_Math_50_Unique_Questions.xlsx', data_only=True)

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    print(f"=== Sheet: {sheetname} ===")
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) > 3:
        h = [str(cell) for cell in rows[3] if cell is not None]
        print("  Headers (row 4):", h)
        # Check first non-empty row
        for r in rows[4:10]:
            if r and any(r):
                non_empty = [(idx, str(val)[:30]) for idx, val in enumerate(r) if val is not None]
                print("  Row sample:", non_empty)
                break

