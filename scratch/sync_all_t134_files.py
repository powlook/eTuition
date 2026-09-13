import sqlite3
import json
import os
import openpyxl

db_paths = [
    'server/etuition.db',
    'QBank/server/qbank.db',
    'QBank/server/etuition.db'
]

json_paths = [
    'questions.json',
    'QBank/questions.json',
    'QBank/public/questions.json',
    'QBank/dist/questions.json'
]

excel_path = 'QBank/questions_bank.xlsx'

# Fetch questions from primary DB
conn = sqlite3.connect('server/etuition.db')
c = conn.cursor()
c.execute("""
    SELECT id, topic_id, question_title, question_text, math_formula, options_json, correct_answer, hint, working_steps_json, image_url, image_alt, difficulty, show_image, show_formula
    FROM questions
    WHERE topic_id = 134 OR question_title LIKE '%T134%'
    ORDER BY id ASC
""")
rows = c.fetchall()
conn.close()

print(f"Fetched {len(rows)} topic 134 questions from primary DB.")

questions = []
for r in rows:
    qid, tid, title, text, formula, opts_json, ans, hint, steps_json, img, alt, diff, show_img, show_form = r
    questions.append({
        'id': qid,
        'topic_id': tid,
        'question_title': title,
        'question_text': text,
        'math_formula': formula,
        'options': json.loads(opts_json) if opts_json else [],
        'correct_answer': ans,
        'hint': hint,
        'working_steps': json.loads(steps_json) if steps_json else [],
        'image_url': img,
        'image_alt': alt,
        'difficulty': diff,
        'show_image': show_img,
        'show_formula': show_form
    })

# Sync SQLite DBs
for db_path in db_paths:
    if not os.path.exists(db_path):
        continue
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    for q in questions:
        c.execute("""
            UPDATE questions
            SET image_url = ?,
                question_title = ?,
                question_text = ?,
                correct_answer = ?
            WHERE id = ?
        """, (q['image_url'], q['question_title'], q['question_text'], q['correct_answer'], q['id']))
    conn.commit()
    conn.close()
    print(f"Synced SQLite DB: {db_path}")

# Sync JSON files
q_map = {q['id']: q for q in questions}

for json_path in json_paths:
    if not os.path.exists(json_path):
        continue
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    if isinstance(data, list):
        for item in data:
            if item.get('id') in q_map:
                u = q_map[item['id']]
                item['image_url'] = u['image_url']
    elif isinstance(data, dict) and 'questions' in data:
        for item in data['questions']:
            if item.get('id') in q_map:
                u = q_map[item['id']]
                item['image_url'] = u['image_url']

    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"Synced JSON file: {json_path}")

# Sync Excel file
if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    id_col = headers.index('id') if 'id' in headers else -1
    
    if id_col != -1:
        for row in ws.iter_rows(min_row=2):
            val = row[id_col].value
            if val in q_map:
                u = q_map[val]
                for col_idx, col_name in enumerate(headers):
                    if col_name == 'image_url':
                        row[col_idx].value = u['image_url']
        wb.save(excel_path)
        print(f"Synced Excel file: {excel_path}")

print("Topic 134 synchronization complete!")
