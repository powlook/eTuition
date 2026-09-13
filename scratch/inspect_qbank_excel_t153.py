import openpyxl

wb = openpyxl.load_workbook('QBank/questions_bank.xlsx')
ws = wb.active

headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
print("QBank/questions_bank.xlsx Headers:", headers)

col_map = {h: idx+1 for idx, h in enumerate(headers) if h}

for r in range(2, ws.max_row+1):
    topic_id_val = str(ws.cell(row=r, column=col_map.get('Topic ID', 2)).value)
    if topic_id_val == '153':
        qid = ws.cell(row=r, column=1).value
        title = ws.cell(row=r, column=col_map.get('Question Title', 3)).value
        text = ws.cell(row=r, column=col_map.get('Question Text', 4)).value
        formula = ws.cell(row=r, column=col_map.get('LaTeX Formula Expression', 5)).value
        opt_a = ws.cell(row=r, column=col_map.get('Option A', 6)).value
        opt_b = ws.cell(row=r, column=col_map.get('Option B', 7)).value
        opt_c = ws.cell(row=r, column=col_map.get('Option C', 8)).value
        opt_d = ws.cell(row=r, column=col_map.get('Option D', 9)).value
        ans = ws.cell(row=r, column=col_map.get('Correct Answer', 10)).value
        
        q_num = r - 1
        t_safe = str(title).encode('ascii', 'ignore').decode('ascii')
        txt_safe = str(text)[:70].encode('ascii', 'ignore').decode('ascii')
        ans_safe = str(ans).encode('ascii', 'ignore').decode('ascii')
        print(f"\nID={qid} | Title={t_safe}")
        print(f"  Text: {txt_safe}")
        print(f"  Ans: {ans_safe}")
        opt_a_s = str(opt_a).encode('ascii', 'ignore').decode('ascii')
        opt_b_s = str(opt_b).encode('ascii', 'ignore').decode('ascii')
        opt_c_s = str(opt_c).encode('ascii', 'ignore').decode('ascii')
        opt_d_s = str(opt_d).encode('ascii', 'ignore').decode('ascii')
        print(f"  Opts: A={opt_a_s} | B={opt_b_s} | C={opt_c_s} | D={opt_d_s}")
