import openpyxl
import json

wb = openpyxl.load_workbook('QBank/resources/Grade_9_Math_50_Unique_Questions.xlsx')
ws = wb['Measurement & Geometry (MG)']

t09_qs = []
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
col_map = {h: idx+1 for idx, h in enumerate(headers) if h}

for r in range(2, ws.max_row+1):
    topic_val = str(ws.cell(row=r, column=col_map.get('Topic Code & Name', 2)).value)
    if 'T09' in topic_val or 'T152' in topic_val:
        q_num = ws.cell(row=r, column=col_map.get('Q#', 1)).value
        q_title = ws.cell(row=r, column=col_map.get('Question Title', 3)).value
        q_text = ws.cell(row=r, column=col_map.get('Question Text', 4)).value
        q_formula = ws.cell(row=r, column=col_map.get('LaTeX Formula Expression', 5)).value
        opt_a = ws.cell(row=r, column=col_map.get('Option A', 6)).value
        opt_b = ws.cell(row=r, column=col_map.get('Option B', 7)).value
        opt_c = ws.cell(row=r, column=col_map.get('Option C', 8)).value
        opt_d = ws.cell(row=r, column=col_map.get('Option D', 9)).value
        answer = ws.cell(row=r, column=col_map.get('Correct Answer', 10)).value
        hint = ws.cell(row=r, column=col_map.get('Hint', 11)).value
        steps = ws.cell(row=r, column=col_map.get('Working Steps', 12)).value
        img_title = ws.cell(row=r, column=col_map.get('Image ALT', 14)).value
        
        t09_qs.append({
            'num': len(t09_qs) + 1,
            'q_num': q_num,
            'title': q_title,
            'text': q_text,
            'formula': q_formula,
            'options': [opt_a, opt_b, opt_c, opt_d],
            'answer': answer,
            'hint': hint,
            'steps': steps,
            'img_title': img_title
        })

print(f"Extracted {len(t09_qs)} questions for T09 / T152.")

with open('scratch/t09_excel_50_questions.json', 'w', encoding='utf-8') as f:
    json.dump(t09_qs, f, indent=2, ensure_ascii=False)

for q in t09_qs:
    t_safe = str(q['title']).encode('ascii', 'ignore').decode('ascii')
    txt_safe = str(q['text'])[:80].encode('ascii', 'ignore').decode('ascii')
    print(f"Q{q['num']}: {t_safe} | Text: {txt_safe}...")
