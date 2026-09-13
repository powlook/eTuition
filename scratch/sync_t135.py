import sqlite3
import json
import os
import openpyxl

# Read questions from server/etuition.db for topic_id = 135
conn_main = sqlite3.connect('server/etuition.db')
cursor_main = conn_main.cursor()
cursor_main.execute("SELECT id, topic_id, question_title, question_text, math_formula, options_json, correct_answer, hint, working_steps_json, image_url, image_alt, difficulty FROM questions WHERE topic_id = 135 ORDER BY id")
main_qs = cursor_main.fetchall()
print(f"Main DB server/etuition.db has {len(main_qs)} questions for Topic 135.")

# 1. Sync to QBank/server/qbank.db and QBank/server/etuition.db
other_dbs = ['QBank/server/qbank.db', 'QBank/server/etuition.db']
for db_path in other_dbs:
    if os.path.exists(db_path):
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        for q in main_qs:
            q_id, t_id, qtitle, qtext, formula, opts_json, ans, hint, steps_json, img_url, img_alt, diff = q
            # Check if row exists
            cursor.execute("SELECT id FROM questions WHERE id = ?", (q_id,))
            if cursor.fetchone():
                cursor.execute("""
                    UPDATE questions SET 
                        question_title = ?, question_text = ?, math_formula = ?,
                        options_json = ?, correct_answer = ?, hint = ?,
                        working_steps_json = ?, image_url = ?, image_alt = ?, difficulty = ?
                    WHERE id = ?
                """, (qtitle, qtext, formula, opts_json, ans, hint, steps_json, img_url, img_alt, diff, q_id))
            else:
                cursor.execute("""
                    INSERT INTO questions (id, topic_id, question_title, question_text, math_formula, options_json, correct_answer, hint, working_steps_json, image_url, image_alt, difficulty)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (q_id, t_id, qtitle, qtext, formula, opts_json, ans, hint, steps_json, img_url, img_alt, diff))
        conn.commit()
        conn.close()
        print(f"Synced {db_path}")

# 2. Sync to JSON files
json_paths = [
    'questions.json',
    'QBank/questions.json',
    'QBank/public/questions.json',
    'QBank/dist/questions.json'
]

for jp in json_paths:
    if os.path.exists(jp):
        with open(jp, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Remove old T135 questions
        data = [q for q in data if str(q.get('topic_id')) != '135']
        
        # Append updated T135 questions
        for q in main_qs:
            q_id, t_id, qtitle, qtext, formula, opts_json, ans, hint, steps_json, img_url, img_alt, diff = q
            data.append({
                'id': q_id,
                'topic_id': t_id,
                'question_title': qtitle,
                'question_text': qtext,
                'math_formula': formula,
                'options_json': opts_json,
                'correct_answer': ans,
                'hint': hint,
                'working_steps_json': steps_json,
                'image_url': img_url,
                'image_alt': img_alt,
                'difficulty': diff
            })
            
        with open(jp, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f"Synced {jp}")

# 3. Sync to Excel bank
excel_path = 'QBank/questions_bank.xlsx'
if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    
    # Map questions by Question ID
    q_map = {q[0]: q for q in main_qs}
    
    for row_idx in range(2, sheet.max_row + 1):
        q_id_cell = sheet.cell(row=row_idx, column=1).value
        if q_id_cell in q_map:
            q = q_map[q_id_cell]
            q_id, t_id, qtitle, qtext, formula, opts_json, ans, hint, steps_json, img_url, img_alt, diff = q
            opts = json.loads(opts_json) if opts_json else []
            
            # Find column indices
            col_map = {h: idx+1 for idx, h in enumerate(headers)}
            if 'Question Title' in col_map: sheet.cell(row=row_idx, column=col_map['Question Title'], value=qtitle)
            if 'Question Text' in col_map: sheet.cell(row=row_idx, column=col_map['Question Text'], value=qtext)
            if 'Option A' in col_map and len(opts)>0: sheet.cell(row=row_idx, column=col_map['Option A'], value=opts[0])
            if 'Option B' in col_map and len(opts)>1: sheet.cell(row=row_idx, column=col_map['Option B'], value=opts[1])
            if 'Option C' in col_map and len(opts)>2: sheet.cell(row=row_idx, column=col_map['Option C'], value=opts[2])
            if 'Option D' in col_map and len(opts)>3: sheet.cell(row=row_idx, column=col_map['Option D'], value=opts[3])
            if 'Correct Answer' in col_map: sheet.cell(row=row_idx, column=col_map['Correct Answer'], value=ans)
            if 'Image URL' in col_map: sheet.cell(row=row_idx, column=col_map['Image URL'], value=img_url)
            
    wb.save(excel_path)
    print(f"Synced {excel_path}")

