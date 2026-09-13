import sqlite3
import json
import openpyxl
import os

with open('scratch/t149_50_perfect_built.json', 'r', encoding='utf-8') as f:
    t149_data = json.load(f)

print(f"Loaded {len(t149_data)} questions from scratch/t149_50_perfect_built.json")

# Map of question id to dict
t149_map = {q['id']: q for q in t149_data}

# 1. Update SQLite Databases
dbs = [
    'server/etuition.db',
    'QBank/server/qbank.db',
    'QBank/server/etuition.db'
]

for db_path in dbs:
    if not os.path.exists(db_path):
        print(f"Database {db_path} does not exist, skipping.")
        continue
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    
    for qid, q in t149_map.items():
        opts = [q['optA'], q['optB'], q['optC'], q['optD']]
        opts_str = json.dumps(opts, ensure_ascii=False)
        
        c.execute("""
            UPDATE questions 
            SET question_title = ?,
                question_text = ?,
                math_formula = ?,
                options_json = ?,
                correct_answer = ?,
                show_image = ?,
                image_url = ?
            WHERE id = ? OR id = ?
        """, (
            q['title'],
            q['question_text'],
            q['math_formula'],
            opts_str,
            q['correct_answer'],
            q['show_image'],
            q['image_url'],
            int(qid),
            f"T149-Q{qid}"
        ))
    conn.commit()
    conn.close()
    print(f"Updated SQLite database: {db_path}")

# 2. Update JSON datastores
jsons = [
    'questions.json',
    'QBank/questions.json',
    'QBank/public/questions.json',
    'QBank/dist/questions.json'
]

for jpath in jsons:
    if not os.path.exists(jpath):
        print(f"JSON path {jpath} does not exist, skipping.")
        continue
    with open(jpath, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    updated_count = 0
    for item in data:
        item_id_str = str(item.get('id') or item.get('question_id') or '')
        num_id = item_id_str.replace('T149-Q', '').replace('Q', '')
        if num_id in t149_map:
            q = t149_map[num_id]
            opts = [q['optA'], q['optB'], q['optC'], q['optD']]
            item['question_title'] = q['title']
            item['question_text'] = q['question_text']
            item['math_formula'] = q['math_formula']
            item['formula'] = q['math_formula']
            item['options'] = opts
            item['options_json'] = opts
            item['correct_answer'] = q['correct_answer']
            item['show_image'] = q['show_image']
            item['image_url'] = q['image_url']
            updated_count += 1

    with open(jpath, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    print(f"Updated JSON datastore ({updated_count} items): {jpath}")

# 3. Update Excel file
excel_path = 'QBank/questions_bank.xlsx'
if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    header_map = {h: idx + 1 for idx, h in enumerate(headers) if h}
    
    updated_rows = 0
    for row in range(2, ws.max_row + 1):
        q_id_val = str(ws.cell(row=row, column=header_map.get('id', 1)).value or '')
        num_id = q_id_val.replace('T149-Q', '').replace('Q', '')
        if num_id in t149_map:
            q = t149_map[num_id]
            if 'question_title' in header_map:
                ws.cell(row=row, column=header_map['question_title'], value=q['title'])
            if 'question_text' in header_map:
                ws.cell(row=row, column=header_map['question_text'], value=q['question_text'])
            if 'math_formula' in header_map:
                ws.cell(row=row, column=header_map['math_formula'], value=q['math_formula'])
            if 'optA' in header_map:
                ws.cell(row=row, column=header_map['optA'], value=q['optA'])
            if 'optB' in header_map:
                ws.cell(row=row, column=header_map['optB'], value=q['optB'])
            if 'optC' in header_map:
                ws.cell(row=row, column=header_map['optC'], value=q['optC'])
            if 'optD' in header_map:
                ws.cell(row=row, column=header_map['optD'], value=q['optD'])
            if 'correct_answer' in header_map:
                ws.cell(row=row, column=header_map['correct_answer'], value=q['correct_answer'])
            if 'show_image' in header_map:
                ws.cell(row=row, column=header_map['show_image'], value=q['show_image'])
            if 'image_url' in header_map:
                ws.cell(row=row, column=header_map['image_url'], value=q['image_url'])
            updated_rows += 1
            
    wb.save(excel_path)
    print(f"Updated Excel datastore ({updated_rows} rows): {excel_path}")

print("T149 synchronization complete!")
