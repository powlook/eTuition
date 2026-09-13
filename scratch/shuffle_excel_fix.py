import openpyxl
import json
import sqlite3
import os

excel_path = "QBank/questions_bank.xlsx"

if not os.path.exists(excel_path):
    print("Excel file not found!")
    exit(0)

wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

# Column mapping from printed output:
# Col 5: Topic ID
# Col 12: Option A
# Col 13: Option B
# Col 14: Option C
# Col 15: Option D
# Col 16: Correct Answer

t_col = 5
opta_c = 12
optb_c = 13
optc_c = 14
optd_c = 15
corr_c = 16

topics = [144, 148, 149, 150, 152, 159]

for t in topics:
    conn = sqlite3.connect("server/etuition.db")
    cursor = conn.cursor()
    cursor.execute("SELECT id, options_json, correct_answer FROM questions WHERE topic_id = ? ORDER BY id ASC", (t,))
    db_rows = cursor.fetchall()
    conn.close()
    
    matching_rows = []
    for r in range(2, sheet.max_row + 1):
        cell_val = sheet.cell(row=r, column=t_col).value
        if cell_val is not None:
            try:
                if int(cell_val) == t:
                    matching_rows.append(r)
            except ValueError:
                pass
    
    matching_rows.sort()
    print(f"Topic {t}: Found {len(matching_rows)} matching rows in Excel (expected {len(db_rows)})")
    
    for idx, r_data in enumerate(db_rows):
        if idx < len(matching_rows):
            row_num = matching_rows[idx]
            opts = json.loads(r_data[1]) if r_data[1] else []
            correct = r_data[2]
            
            sheet.cell(row=row_num, column=opta_c, value=opts[0] if len(opts)>0 else "")
            sheet.cell(row=row_num, column=optb_c, value=opts[1] if len(opts)>1 else "")
            sheet.cell(row=row_num, column=optc_c, value=opts[2] if len(opts)>2 else "")
            sheet.cell(row=row_num, column=optd_c, value=opts[3] if len(opts)>3 else "")
            sheet.cell(row=row_num, column=corr_c, value=correct)

wb.save(excel_path)
print("Successfully updated Excel bank for all topics!")
