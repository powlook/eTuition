import openpyxl

wb = openpyxl.load_workbook('QBank/resources/Grade_8_Math_50_Unique_Questions.xlsx', data_only=True)
for sname in wb.sheetnames:
    sheet = wb[sname]
    print(f"=== Sheet {sname} ===")
    for row in list(sheet.iter_rows(values_only=True))[:5]:
        print(" ", [str(c)[:30] if c is not None else '' for c in row[:10]])

