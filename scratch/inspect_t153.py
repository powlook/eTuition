import sqlite3
import openpyxl
import json

print("=== 1. SERVER/ETUITION.DB FOR TOPIC 153 ===")
conn = sqlite3.connect('server/etuition.db')
c = conn.cursor()

# Get topic details
c.execute("SELECT * FROM topics WHERE id = 153")
t_row = c.fetchone()
print(f"Topic 153 Details: {t_row}")

# Get all questions for Topic 153
c.execute("SELECT id, question_title, question_text, math_formula, correct_answer, image_url, show_image FROM questions WHERE topic_id = 153 ORDER BY id")
rows = c.fetchall()
print(f"Total questions for Topic 153 in server/etuition.db: {len(rows)}")

for idx, r in enumerate(rows[:15]):
    t_safe = r[1].encode('ascii', 'ignore').decode('ascii') if r[1] else ''
    txt_safe = r[2][:80].encode('ascii', 'ignore').decode('ascii') if r[2] else ''
    print(f"Q{idx+1} (ID {r[0]}): Title='{t_safe}' | Img={r[5]} | ShowImg={r[6]} | Text='{txt_safe}'")

conn.close()

print("\n=== 2. QBANK/RESOURCES EXCEL FOR T10 (THE TRIGONOMETRIC RATIOS) ===")
wb = openpyxl.load_workbook('QBank/resources/Grade_9_Math_50_Unique_Questions.xlsx')
ws = wb['Measurement & Geometry (MG)']

headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
col_map = {h: idx+1 for idx, h in enumerate(headers) if h}

t10_qs = []
for r in range(2, ws.max_row+1):
    topic_val = str(ws.cell(row=r, column=col_map.get('Topic Code & Name', 2)).value)
    if 'T10' in topic_val or '153' in topic_val:
        q_num = ws.cell(row=r, column=col_map.get('Q#', 1)).value
        q_title = ws.cell(row=r, column=col_map.get('Question Title', 3)).value
        q_text = ws.cell(row=r, column=col_map.get('Question Text', 4)).value
        t10_qs.append((q_num, q_title, q_text))

print(f"Total T10 questions in Excel: {len(t10_qs)}")
for q in t10_qs[:10]:
    t_safe = str(q[1]).encode('ascii', 'ignore').decode('ascii')
    txt_safe = str(q[2])[:80].encode('ascii', 'ignore').decode('ascii')
    print(f"Excel T10 Q{q[0]}: Title='{t_safe}' | Text='{txt_safe}'")
