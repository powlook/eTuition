import openpyxl
import os

excel_paths = [
    'QBank/resources/Grade_8_Math_50_Unique_Questions.xlsx',
    'QBank/resources/Grade_8_Math_Questions.xlsx'
]

for ep in excel_paths:
    if os.path.exists(ep):
        wb = openpyxl.load_workbook(ep, read_only=True)
        print(f"=== {ep} ===")
        print("Sheet names:", wb.sheetnames)
        for sname in wb.sheetnames:
            if '134' in sname or 'System' in sname or 'Linear' in sname:
                sheet = wb[sname]
                print(f"  Sheet '{sname}': {sheet.max_row} rows")

